"""CSV/XLSX business candidates, resolved against the current Corpus only."""
import csv
import hashlib
import io
import re
from itertools import islice
from zipfile import BadZipFile
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError
from .governance import _answer_anchor_supported, _normalized

COLUMNS = ['Question', 'Reference Answer', 'Evidence', 'Document', 'Question Type', 'Product', 'Version', 'Notes', 'Evaluation Group', 'Expected Behavior', 'Negative Subtype', 'Ablation Attribute', 'Original Entity', 'Alias Expression']
GROUPS = {'positive', 'ablation', 'negative'}
BEHAVIORS = {'clarify', 'insufficient_evidence', 'safe_rejection', 'prompt_injection_resistance'}
SUBTYPES = {'clarify', 'insufficient_evidence', 'safe_rejection', 'safety_critical', 'prompt_injection'}

def construction_errors(kind, answer, sources):
    if not kind:
        return []
    texts = [item.get('chunk_text', item.get('text', '')) for item in sources]
    facts = [part.strip() for text in texts for part in text.splitlines() if re.search(r'[^：:]+[：:]\s*\S+', part)]
    if kind == 'Fact':
        return [] if any(_answer_anchor_supported(answer, [fact]) for fact in facts) else ['Fact 需要可定位的实体属性值']
    if kind == 'Aggregation':
        return [] if len(facts) >= 2 and all(_answer_anchor_supported(fact, [answer]) for fact in facts) else ['Aggregation 需要完整属性材料及穷举答案']
    if kind == 'Bridge':
        entities = [set(re.findall(r'[A-Za-z]+\d+|[\u4e00-\u9fffA-Za-z]+[A-Za-z0-9]+', text)) for text in texts]
        return [] if len(texts) >= 2 and set.intersection(*entities) and all(_answer_anchor_supported(answer, [text]) is False for text in texts) and _answer_anchor_supported(answer, texts) else ['Bridge 需要共享实体与两段共同支撑的证据']
    return ['Question Type 仅支持空值、Fact、Aggregation、Bridge']

def template(format):
    if format == 'csv':
        output = io.StringIO(); csv.writer(output).writerow(COLUMNS)
        return ('\ufeff' + output.getvalue()).encode('utf-8'), 'text/csv'
    if format == 'xlsx':
        from openpyxl import Workbook
        workbook = Workbook(); workbook.active.append(COLUMNS)
        workbook.active.freeze_panes = 'A2'
        guide = workbook.create_sheet('填写说明')
        guide.append(['正向/消融：Evidence 填当前 Corpus 的逐字原文，多个证据用换行分隔；Document 填文档名或 ID，多个用分号分隔。'])
        guide.append(['Evaluation Group 必填 positive/ablation/negative；构造题型可留空。消融题填写 Ablation Attribute。'])
        guide.append(['Negative 填 Expected Behavior 与 Negative Subtype，留空答案和证据。导入不会执行模型或审批。'])
        output = io.BytesIO(); workbook.save(output)
        return output.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    raise ValueError('仅支持 CSV / XLSX')

def parse_import(data, filename, chunks, existing):
    if len(data) > 10 * 1024 * 1024:
        raise ValueError('文件超过 10 MB')
    if filename.lower().endswith('.csv'):
        try:
            rows = list(islice(csv.reader(io.StringIO(data.decode('utf-8-sig')), strict=True), 1002))
        except csv.Error as error:
            raise ValueError('CSV 格式无效') from error
    elif filename.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        from zipfile import ZipFile
        try:
            with ZipFile(io.BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError('XLSX 解压内容超过限制')
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        except (BadZipFile, InvalidFileException, ParseError, KeyError) as error:
            raise ValueError('XLSX 格式无效') from error
        try:
            rows = list(islice(workbook.active.iter_rows(values_only=True), 1002))
        finally:
            workbook.close()
    else:
        raise ValueError('仅支持 CSV / XLSX')
    if not rows or len(rows) > 1001:
        raise ValueError('需要表头，且最多 1000 行')
    headers = [str(value or '').strip() for value in rows[0]]
    if len(headers) != len(set(headers)) or not {'Question', 'Evaluation Group'}.issubset(headers):
        raise ValueError('表头重复或缺少 Question / Evaluation Group')
    seen = {_normalized(item['question']) for item in existing}
    previews, valid = [], []
    for index, row in enumerate(rows[1:], 2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        fields = {key: str(value or '').strip() for key, value in zip(headers, row)}
        errors = []
        if any(value.startswith('=') for value in fields.values()): errors.append('不接受公式单元格')
        if len(row) > len(headers): errors.append('数据列数超过表头')
        question, group = fields.get('Question', ''), fields.get('Evaluation Group', '').lower()
        key = _normalized(question)
        if not question or len(question) > 1000: errors.append('Question 必填且最多 1000 字')
        if key in seen: errors.append('重复题目')
        seen.add(key)
        if group not in GROUPS: errors.append('Evaluation Group 必须显式填写 positive / ablation / negative')
        answer, quote = fields.get('Reference Answer', ''), fields.get('Evidence', '')
        documents = set(re.split(r'[;；\n]', fields.get('Document', '')))
        sources = []
        if group in {'positive', 'ablation'}:
            for fragment in quote.splitlines():
                matches = [chunk for chunk in chunks if fragment.strip() and fragment.strip() in chunk.get('chunk_text', chunk.get('text', '')) and (chunk.get('document_id') in documents or chunk.get('document_name') in documents)]
                if not matches: errors.append('Evidence 无法定位到指定的当前文档原文')
                sources.extend(matches)
            sources = list({item['chunk_id']: item for item in sources}.values())
            if not answer or len(answer) > 4000 or not sources: errors.append('正向/消融需要答案与可定位证据；文档名称不能替代证据')
            elif not _answer_anchor_supported(answer, [item.get('chunk_text', item.get('text', '')) for item in sources]): errors.append('答案缺少原文支持')
            errors.extend(construction_errors(fields.get('Question Type'), answer, sources))
        if group == 'negative':
            if fields.get('Expected Behavior') not in BEHAVIORS or fields.get('Negative Subtype') not in SUBTYPES: errors.append('Negative 需要有效 Expected Behavior / Negative Subtype')
            if answer or quote or fields.get('Question Type'): errors.append('Negative 不填写普通答案、证据或构造题型')
        attribute = fields.get('Ablation Attribute')
        if group == 'ablation':
            if attribute not in {'cross_chunk', 'alias_entity', 'weak_keywords', 'colloquial'}: errors.append('消融题需填写 cross_chunk / alias_entity / weak_keywords / colloquial')
            if attribute == 'cross_chunk' and len(sources) < 2: errors.append('跨段题需要至少两段证据')
            if attribute == 'alias_entity' and not all(fields.get(name) for name in ('Original Entity', 'Alias Expression')): errors.append('别名题需要原始实体与别名')
        candidate = {'question': question, 'reference_answer': answer or None, 'test_category': group, 'negative_subtype': fields.get('Negative Subtype') or None, 'expected_behavior': fields.get('Expected Behavior') or None, 'construction_type': fields.get('Question Type') or None, 'ablation_attribute': attribute or None, 'ablation_metadata': {'original_entity': fields.get('Original Entity'), 'alias_expression': fields.get('Alias Expression')}, 'evidence': [{'source_chunk_ids': [item['chunk_id'] for item in sources], 'evidence_key_points': quote.splitlines()}] if sources else [], 'import_fields': fields, 'import_row': index}
        previews.append({'row': index, 'fields': fields, 'errors': errors})
        if not errors: valid.append(candidate)
    return {'fields': headers, 'rows': previews, 'valid_rows': valid, 'valid_count': len(valid), 'error_count': sum(bool(row['errors']) for row in previews), 'file_hash': hashlib.sha256(data).hexdigest()}
