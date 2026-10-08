"""Candidate pool closure: isolated storage, no external providers."""
import io
import csv
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from app.business_import import BUSINESS_COLUMNS, COLUMNS, template, parse_import
from app.golden_v2 import match_pool
from app.governance import GovernanceStore


class BusinessClosureTests(unittest.TestCase):
    def workbook(self, rows, headers=BUSINESS_COLUMNS):
        book = Workbook(); book.active.title = '业务问题'; book.active.append(headers)
        for row in rows: book.active.append(row)
        example = book.create_sheet('填写示例'); example.append(['不应导入']); book.active = 1
        data = io.BytesIO(); book.save(data); return data.getvalue()

    def test_default_excel_four_columns_style_and_separate_examples(self):
        data, _ = template('xlsx'); book = load_workbook(io.BytesIO(data))
        self.assertEqual(book.sheetnames, ['业务问题', '填写示例', '填写说明'])
        sheet = book['业务问题']; self.assertEqual([c.value for c in sheet[1]], BUSINESS_COLUMNS)
        self.assertEqual(sheet.max_row, 1); self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet['A1'].fill.fgColor.rgb, '00EEEEEE')
        self.assertTrue(sheet['A1'].alignment.wrap_text)
        self.assertTrue(all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row))
        self.assertEqual(parse_import(data, 'template.xlsx', [], [])['recognized_count'], 0)

    def test_chinese_context_not_verified_answer_or_fake_evidence(self):
        data = self.workbook([['设备报警该怎么办？', 'CR10', '现场维护', '先断电']])
        result = parse_import(data, 'business.xlsx', [], [])
        self.assertEqual(result['valid_count'], 1); self.assertEqual(result['error_count'], 0)
        candidate = result['valid_rows'][0]
        self.assertEqual(candidate['import_context'], {'scenario': '现场维护', 'provided_reference': '先断电'})
        self.assertIsNone(candidate['reference_answer']); self.assertEqual(candidate['evidence'], [])
        self.assertEqual(candidate['test_category'], 'unclassified'); self.assertFalse(candidate['validation']['valid'])

    def test_partial_rows_errors_line_numbers_duplicate_formula_blank(self):
        data = self.workbook([['设备报警该怎么办？'], ['设备报警该怎么办？'], ['', 'CR10'], ['=1+1']])
        result = parse_import(data, 'business.xlsx', [], [])
        self.assertEqual((result['valid_count'], result['error_count']), (1, 3))
        self.assertEqual([r['row'] for r in result['rows'] if r['errors']], [3, 4, 5])
        self.assertIn('不接受公式单元格', result['rows'][-1]['errors'])

    def test_legacy_csv_and_xlsx_keep_validation(self):
        fields = {'Question': '机器人需要先断电吗？', 'Reference Answer': '先断电', 'Evidence': '先断电', 'Evaluation Group': 'positive'}
        row = [fields.get(name, '') for name in COLUMNS]
        text = io.StringIO(); writer = csv.writer(text); writer.writerow(COLUMNS); writer.writerow(row)
        chunks = [{'chunk_id': 'C1', 'document_id': 'D1', 'chunk_text': '维修前先断电。'}]
        for data, name in [(text.getvalue().encode(), 'old.csv'), (self.workbook([row], COLUMNS), 'old.xlsx')]:
            result = parse_import(data, name, chunks, [])
            self.assertEqual(result['format'], 'legacy'); self.assertEqual(result['valid_count'], 1)
            self.assertEqual(result['valid_rows'][0]['reference_answer'], '先断电')

    def test_transaction_receipt_replay_and_pending_qualification(self):
        with tempfile.TemporaryDirectory() as directory:
            store = GovernanceStore(Path(directory) / 'test.db')
            result = parse_import(self.workbook([['设备如何停机？']]), 'b.xlsx', [], [])
            ids = store.save_business_candidates(result['valid_rows'], 'b.xlsx', result['file_hash'], [{'row': 3, 'errors': ['重复题目']}])
            self.assertEqual(ids, store.save_business_candidates(result['valid_rows'], 'b.xlsx', result['file_hash']))
            self.assertEqual(len(store.questions()), 41)
            self.assertEqual(store.business_import_receipt(result['file_hash'])['errors'][0]['row'], 3)
            self.assertEqual(store.candidate_rows(ids)[0]['qualification_status'], 'needs_human_review')
            self.assertEqual(store.question(ids[0])['raw']['source'], 'business_import')
            with self.assertRaises(ValueError):
                store.save_business_candidates([{**result['valid_rows'][0], 'question': '另一个问题'}, result['valid_rows'][0]], 'b.xlsx', 'other')
            self.assertEqual(len(store.questions()), 41); self.assertIsNone(store.business_import_receipt('other'))

    def test_active_list_collapses_only_real_lineage_keeps_history(self):
        with tempfile.TemporaryDirectory() as directory:
            store = GovernanceStore(Path(directory) / 'test.db')
            candidate = parse_import(self.workbook([['原始问题']]), 'b.xlsx', [], [])['valid_rows'][0]
            old = store.save_business_candidates([candidate], 'b.xlsx', 'one')[0]
            new = store.save_business_candidates([{**candidate, 'question': '修订问题', 'replaces_question_id': old}], 'c.xlsx', 'two')[0]
            self.assertEqual([r['id'] for r in store.active_pool_rows()], [new])
            self.assertEqual(len(store.questions()), 42)

    def test_max_matching_preserves_preferred_selection_and_reports_gaps(self):
        plan = {'plan_id': 'p', 'profile': {'positive_count': 2, 'ablation_count': 0, 'negative_count': 0}, 'slots': [{'slot_id': str(i), 'evaluation_group': 'positive', 'topic_cluster': 'T', 'construction_type': 'Ordinary'} for i in range(2)]}
        valid = lambda slots: {'valid': True, 'normalized_candidate': {'test_category': 'positive'}, 'coverage_match': {'eligible_slot_ids': slots}}
        matching = match_pool(plan, {'a': valid(['0','1']), 'b': valid(['0']), 'z': valid(['0','1'])}, preferred_ids=['z'])
        self.assertIn('z', matching['matching'].values()); self.assertEqual(len(matching['matching']), 2)
        short = match_pool(plan, {'z': valid(['0'])}, preferred_ids=['z'])
        self.assertFalse(short['valid']); self.assertEqual(len(short['gaps']), 1)

    def test_api_partial_requires_explicit_confirmation_replay_and_no_provider(self):
        from api_fixture import main
        from fastapi.testclient import TestClient
        with tempfile.TemporaryDirectory() as directory:
            store = GovernanceStore(Path(directory) / 'test.db')
            data = self.workbook([['如何安全停机？'], ['如何安全停机？']])
            with patch.object(main, 'store', store), patch.object(main, 'current_coverage_plan', return_value=None), patch.object(main.corpus, 'chunks', return_value=[]), patch.object(main.ai_service, 'negative_topic_embedding', side_effect=AssertionError('No Provider')):
                client = TestClient(main.app)
                url = '/api/governance/imports?filename=business.xlsx'
                preview = client.post(url, content=data).json()
                denied = client.post(url + '&confirm=true', content=data)
                self.assertEqual(denied.status_code, 422); self.assertEqual(len(store.questions()), 40)
                confirmed = client.post(url + '&confirm=true&allow_partial=true&preview_hash=' + preview['file_hash'], content=data)
                self.assertEqual(confirmed.status_code, 200); self.assertEqual(confirmed.json()['imported_count'], 1)
                replay = client.post(url + '&confirm=true&allow_partial=true', content=data)
                self.assertTrue(replay.json()['replayed']); self.assertEqual(replay.json()['created_count'], 0)
                self.assertEqual(len(store.questions()), 41)
                changed = client.post(url + '&confirm=true&preview_hash=wrong', content=data)
                self.assertEqual(changed.status_code, 422)

    def test_complete_medium_fifty_creates_only_unapproved_working_set(self):
        import numpy as np
        from app.governance import GENERATION_PROFILES
        with tempfile.TemporaryDirectory() as directory:
            store = GovernanceStore(Path(directory) / 'test.db')
            chunks = [{'chunk_id': f'C{i:03}', 'document_id': 'D', 'chunk_text': '断电后维护设备。'} for i in range(160)]
            plan = store.coverage_preview('medium', chunks, np.repeat(np.eye(4, dtype='float32'), 40, axis=0))
            candidates = []
            for index, slot in enumerate(plan['slots']):
                negative = slot['evaluation_group'] == 'negative'
                candidates.append({'question': f'隔离测试场景{index}如何维护设备？', 'test_category': slot['evaluation_group'], 'construction_type': 'Ordinary', 'reference_answer': None if negative else '断电后维护设备。', 'evidence': [] if negative else [{'source_chunk_ids': slot['material_chunk_ids']}], 'ablation_attribute': slot['ablation_attribute'], 'expected_behavior': slot.get('expected_behavior'), 'negative_subtype': slot.get('negative_subtype'), 'generation_strategy': 'business_v2' if negative else 'offline_fixture', 'coverage_anchor_chunk_ids': slot['material_chunk_ids'] if negative else [], 'import_row': index+2})
            ids = store.save_business_candidates(candidates, 'offline-fixture.xlsx', 'fixture')
            before_snapshots = store.dataset_snapshots()
            preview = store.preview_pool_run('medium', ids, chunks, plan['plan_id'])
            self.assertTrue(preview['valid'], preview['gaps']); self.assertEqual(preview['counts'], {'positive':20,'ablation':10,'negative':20})
            with patch.object(store, 'active_pool_rows', return_value=[{'id': key, 'qualification_status': 'machine_qualified'} for key in ids]):
                filled = store.autofill_pool_run('medium', ids[:2], chunks, plan['plan_id'])
            self.assertTrue(filled['valid']); self.assertEqual(len(filled['question_ids']), 50)
            self.assertTrue(set(ids[:2]).issubset(filled['question_ids']))
            with patch.object(store, 'active_pool_rows', return_value=[]):
                shortage = store.autofill_pool_run('medium', ids[:2], chunks, plan['plan_id'])
            self.assertFalse(shortage['valid']); self.assertEqual(shortage['question_ids'], ids[:2])
            from api_fixture import main
            from fastapi.testclient import TestClient
            with patch.object(main, 'store', store), patch.object(main, 'current_coverage_plan', return_value=plan), patch.object(main.corpus, 'chunks', return_value=chunks), patch.object(main.ai_service, 'negative_topic_embedding', side_effect=AssertionError('No Provider')):
                response = TestClient(main.app).post('/api/governance/generation-runs/from-pool', json={'profile': 'medium', 'question_ids': ids, 'plan_id': plan['plan_id']})
            self.assertEqual(response.status_code, 201, response.text)
            run = response.json()
            self.assertEqual(len(run['question_ids']), 50)
            self.assertTrue(all(row['probe_status']=='probe_pending' and row['qc_status']=='qc_pending' and row['review_status']=='human_review_pending' for row in store.candidate_rows(run['question_ids'])))
            self.assertEqual(store.dataset_snapshots(), before_snapshots)
            self.assertFalse(store.preview_pool_run('medium', ids[:2], chunks, plan['plan_id'])['valid'])
